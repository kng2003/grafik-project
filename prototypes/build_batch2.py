# -*- coding: utf-8 -*-
"""
Шаг 4 фазы 4: вторая партия кодов ГЭСН на подтверждение заказчику -
коды со средним скором (0.5-0.7). Партия 1 (скор >=0.7) уже закрыта
заказчиком полностью (46/46, включая ручные правки).
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

SRC = "prototypes/results/2026-09-17_gesn_frsn_candidates_v2.xlsx"
OUT = "prototypes/results/2026-09-19_shag4_partiya2_srednie.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Соответствие ГЭСН-ФРСН"]
rows = list(ws.iter_rows(min_row=2, values_only=True))

def numscore(r):
    v = r[7]
    return v if isinstance(v, (int, float)) else None

selected = [r for r in rows if (numscore(r) is not None and 0.5 <= numscore(r) < 0.7)]
selected.sort(key=lambda r: numscore(r), reverse=True)

out = openpyxl.Workbook()
sh = out.active
sh.title = "Партия 2 - средние"

headers = ["№", "Код ГЭСН", "Наименование ГЭСН", "Ед.изм.", "Код ФРСН (кандидат)",
           "Наименование ФРСН", "чел-ч ФРСН", "Состав звена", "Скор",
           "Ваше решение (да / нет / уточнить)", "Комментарий"]
sh.append(headers)
for c in sh[1]:
    c.font = Font(bold=True)
    c.alignment = Alignment(wrap_text=True, vertical="top")

yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

for i, r in enumerate(selected, start=1):
    row = [i, r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], "", ""]
    sh.append(row)

widths = [5, 18, 42, 9, 18, 42, 10, 30, 8, 22, 30]
for col, w in zip("ABCDEFGHIJK", widths):
    sh.column_dimensions[col].width = w

for row_idx in range(2, sh.max_row + 1):
    sh.cell(row_idx, 10).fill = yellow

sh.freeze_panes = "B2"

sh2 = out.create_sheet("Как заполнить")
instructions = [
    "ЧТО ЭТО ЗА ФАЙЛ",
    "Это вторая партия из 4. В неё вошли коды ГЭСН со скором совпадения",
    "0.5-0.7 - средним по надёжности (0.7 и выше - это партия 1, уже закрыта).",
    "Совпадение здесь по смыслу обычно верное, но чаще требует внимания:",
    "могут отличаться детали (диаметр, вес, материал), которых алгоритм не видит",
    "или видит не полностью.",
    "",
    "ЧТО НУЖНО СДЕЛАТЬ",
    "То же самое, что и в партии 1: посмотреть наименование ГЭСН, наименование",
    "ФРСН и состав звена. Если по смыслу совпадает - 'да'. Если явно не подходит -",
    "'нет'. Если не уверены - 'уточнить' и короткий комментарий, что смущает.",
    "",
    "НА ЧТО ОБРАТИТЬ ОСОБОЕ ВНИМАНИЕ В ЭТОЙ ПАРТИИ",
    "Средний скор часто получается из-за того, что в названии ГЭСН и ФРСН",
    "есть числовой параметр (диаметр, вес, объём) - если цифры не совпадают,",
    "это повод написать 'уточнить', даже если по смыслу работа похожая.",
    "",
    "СКОЛЬКО ВРЕМЕНИ ЭТО ЗАЙМЁТ",
    "65 строк, прикидочно 30-40 минут, потому что придётся чаще сверяться",
    "с цифрами в названии, чем в партии 1.",
    "",
    "ЧТО ДАЛЬШЕ",
    "После этой партии - партия 3 (скор 0.3-0.5, 45 кодов), и отдельно",
    "код ГЭСН08-01-008-05, который алгоритм вообще не смог сопоставить.",
]
for line in instructions:
    sh2.append([line])
sh2.column_dimensions['A'].width = 95
for row_idx in range(1, len(instructions) + 1):
    sh2.cell(row_idx, 1).alignment = Alignment(wrap_text=True)
sh2['A1'].font = Font(bold=True, size=13)

out.save(OUT)
print(f"Сохранено: {OUT}")
print(f"Строк в партии 2: {len(selected)}")
