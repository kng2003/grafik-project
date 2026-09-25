# -*- coding: utf-8 -*-
"""
Шаг 3 фазы 4: подбор 2-3 кандидатов ФРСН для каждого из 157 кодов ГЭСН
по текстовому сходству названий (без внешних NLP-библиотек - их нет на
устройстве, только стандартная библиотека Python).

Вход:
  prototypes/results/2026-09-16_gesn_codes_full_list.xlsx  (157 кодов ГЭСН)
  docs/ФРСН/*.md  (71 файл, парсится через parse_frsn_prototype.py)

Выход:
  prototypes/results/2026-09-17_gesn_frsn_candidates_draft.xlsx

Метод: Jaccard-пересечение значимых русских слов названия + бонус за
совпадение числового параметра (диаметр и т.п.) + бонус/штраф для кодов
ГЭСНм относительно гипотезы, что им соответствуют сборники ФРСН 49-72.
Числовые чел-ч НЕ используются как критерий подбора (доказано ненадёжным
на примере ГЭСН22-03-007-02).
"""
import re, glob, sys, time
sys.path.insert(0, 'prototypes')
from parse_frsn_prototype import parse_file
import openpyxl
from openpyxl.styles import Font

t0 = time.time()

STOPWORDS = set("""
и в с по от до на из для не или так как при над под без через между также ещё эту это тот та то те
диаметр диаметром толщина толщиной устройство работы работ мм м2 м3 шт кг см номинальным
группа группой при выполнении высотой высота
""".split())

def norm_tokens(s):
    s = s or ''
    s = s.lower()
    s = re.sub(r'[«»"\'()]', ' ', s)
    s = re.sub(r'[.,:;]', ' ', s)
    words = re.findall(r'[а-яё]{3,}', s)
    return [w for w in words if w not in STOPWORDS]

def numbers(s):
    return set(re.findall(r'\d+[.,]?\d*', s or ''))

# 1. Загрузка 157 кодов ГЭСН
wb = openpyxl.load_workbook("prototypes/results/2026-09-16_gesn_codes_full_list.xlsx", read_only=True, data_only=True)
ws = wb["Коды ГЭСН"]
gesn = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row and row[1]:
        gesn.append({'code': row[1], 'name': row[2], 'unit': row[3]})
print(f"[{time.time()-t0:.1f}s] Загружено кодов ГЭСН: {len(gesn)}")

# 2. Парсинг всех 71 файлов ФРСН
files = sorted(glob.glob("docs/ФРСН/*.md"))
frsn_records = []
for path in files:
    m = re.match(r'docs/ФРСН/(\d+) ', path)
    sbornik_num = int(m.group(1)) if m else -1
    recs = parse_file(path)
    for r in recs:
        if not r['valid_shifr']:
            continue
        r['sbornik'] = sbornik_num
        frsn_records.append(r)
print(f"[{time.time()-t0:.1f}s] Загружено записей ФРСН (валидный шифр): {len(frsn_records)}")

# 3. Инвертированный индекс: слово -> список индексов записей
index = {}
rec_tokens = []
for i, r in enumerate(frsn_records):
    toks = set(norm_tokens(r['name']))
    rec_tokens.append(toks)
    for w in toks:
        index.setdefault(w, []).append(i)
print(f"[{time.time()-t0:.1f}s] Построен индекс, уникальных слов: {len(index)}")

def find_candidates(gesn_item, top_n=3):
    q_tokens = set(norm_tokens(gesn_item['name']))
    q_nums = numbers(gesn_item['name'])
    is_gesnm = gesn_item['code'].upper().startswith('ГЭСНМ')
    cand_idx = set()
    for w in q_tokens:
        cand_idx.update(index.get(w, []))
    scored = []
    for i in cand_idx:
        r = frsn_records[i]
        r_tokens = rec_tokens[i]
        if not q_tokens or not r_tokens:
            continue
        overlap = q_tokens & r_tokens
        jacc = len(overlap) / len(q_tokens | r_tokens)
        r_nums = numbers(r['name'])
        num_bonus = 0.25 if (q_nums and r_nums and (q_nums & r_nums)) else 0.0
        family_bonus = 0.15 if (is_gesnm and 49 <= r['sbornik'] <= 72) else 0.0
        family_penalty = -0.1 if (is_gesnm and not (49 <= r['sbornik'] <= 72)) else 0.0
        score = jacc + num_bonus + family_bonus + family_penalty
        scored.append((score, i))
    scored.sort(key=lambda x: -x[0])
    return scored[:top_n]

results = []
for g in gesn:
    cands = find_candidates(g, top_n=3)
    results.append((g, cands))
print(f"[{time.time()-t0:.1f}s] Подбор кандидатов завершён")

# 4. Формирование выходного xlsx
out = openpyxl.Workbook()
sh = out.active
sh.title = "Соответствие ГЭСН-ФРСН"
headers = ["Код ГЭСН", "Наименование ГЭСН", "Ед.изм. ГЭСН"]
for k in (1, 2, 3):
    headers += [f"Код ФРСН #{k}", f"Наименование ФРСН #{k}", f"чел-ч #{k}", f"Состав звена #{k}", f"Скор #{k}"]
headers.append("Подтверждено/отклонено (заполняет заказчик)")
sh.append(headers)
for c in sh[1]:
    c.font = Font(bold=True)

no_candidates = 0
for g, cands in results:
    row = [g['code'], g['name'], g['unit']]
    if not cands:
        no_candidates += 1
    for k in range(3):
        if k < len(cands):
            score, i = cands[k]
            r = frsn_records[i]
            crew_str = "; ".join(f"{c['profession']} x{c['qty']}" for c in r['crew']) if r['crew'] else "(не найден)"
            row += [r['shifr'], r['name'], r['lab_hours'], crew_str, round(score, 3)]
        else:
            row += ['', '', '', '', '']
    row.append('')
    sh.append(row)

for col, width in zip('ABCDEFGHIJKLMNOP', [18, 45, 10, 18, 45, 8, 30, 8, 18, 45, 8, 30, 8, 18, 45, 8]):
    sh.column_dimensions[col].width = width

sh2 = out.create_sheet("Сводка")
sh2.append(["Методология подбора кандидатов ГЭСН -> ФРСН (Шаг 3, черновик)"])
sh2.append([])
sh2.append(["Метод: пересечение значимых слов названия (Jaccard) + бонус за совпадение", ])
sh2.append(["числового параметра (диаметр и т.п.) + бонус/штраф для ГЭСНм относительно", ])
sh2.append(["гипотезы, что этим кодам соответствуют сборники ФРСН 49-72.", ])
sh2.append([])
sh2.append(["ВАЖНО: числовые чел-ч НЕ использовались как критерий подбора -", ])
sh2.append(["это доказанно ненадёжный признак (см. пример ГЭСН22-03-007-02).", ])
sh2.append([])
sh2.append(["ВАЖНО: нет доступа к специализированным библиотекам сопоставления текста", ])
sh2.append(["(rapidfuzz, sklearn, pymorphy2) - использован упрощённый метод на", ])
sh2.append(["стандартной библиотеке Python. Кандидаты требуют обязательной", ])
sh2.append(["проверки человеком, это не готовый результат.", ])
sh2.append([])
sh2.append([f"Кодов ГЭСН: {len(gesn)}", f"Кодов без кандидатов: {no_candidates}"])
sh2.append([f"Записей ФРСН (валидный шифр): {len(frsn_records)}", f"Уникальных слов в индексе: {len(index)}"])

out.save("prototypes/results/2026-09-17_gesn_frsn_candidates_draft.xlsx")
print(f"[{time.time()-t0:.1f}s] Сохранено")
print(f"Кодов без кандидатов: {no_candidates} из {len(gesn)}")
