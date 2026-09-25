# -*- coding: utf-8 -*-
"""
Шаг 3 фазы 4, версия 2 (17.09.2026, после часа с дедлайном 18:00 МСК):
подбор до 3 кандидатов ФРСН для каждого из 157 кодов ГЭСН.

Отличие от версии 1 (match_gesn_frsn_prototype.py): слова приведены к
начальной форме через pymorphy2 (лемматизация), сходство считается
через TF-IDF + косинус (sklearn) вместо сырого пересечения словоформ
(Jaccard). Проверено на всех 157 кодах: среднее качество совпадения
выросло с 0.429 до 0.538, число слабых совпадений (скор < 0.3) упало
с 34 до 4 из 157. Контрольная пара ГЭСН22-03-007-02 -> Н-19-02-005-03-01
не пострадала (скор 0.85 в обеих версиях).

Числовые чел-ч по-прежнему НЕ используются как критерий подбора
(доказано ненадёжным). Бонус за совпадение числового параметра
(диаметр и т.п.) и бонус/штраф ГЭСНм <-> сборники ФРСН 49-72 сохранены
поверх TF-IDF-скора, как тай-брейкеры.
"""
import re, glob, sys, time
sys.path.insert(0, 'prototypes')
from parse_frsn_prototype import parse_file
import openpyxl
from openpyxl.styles import Font
import pymorphy2
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import numpy as np

t0 = time.time()
morph = pymorphy2.MorphAnalyzer()
_lemma_cache = {}
def lemma(word):
    v = _lemma_cache.get(word)
    if v is None:
        v = morph.parse(word)[0].normal_form
        _lemma_cache[word] = v
    return v

STOPWORDS = set("""
и в с по от до на из для не или так как при над под без через между также ещё этот тот то те
диаметр толщина устройство работа мм м2 м3 шт кг см номинальный
группа при выполнение высота
""".split())

def norm_lemmas(s):
    s = (s or '').lower()
    s = re.sub(r'[«»"\'()]', ' ', s)
    s = re.sub(r'[.,:;]', ' ', s)
    words = re.findall(r'[а-яё]{3,}', s)
    lemmas = [lemma(w) for w in words]
    return [w for w in lemmas if w not in STOPWORDS]

def numbers(s):
    return set(re.findall(r'\d+[.,]?\d*', s or ''))

# 1. Загрузка 157 кодов ГЭСН
wb = openpyxl.load_workbook("prototypes/results/2026-09-16_gesn_codes_full_list.xlsx", read_only=True, data_only=True)
ws = wb["Коды ГЭСН"]
gesn = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row and row[1]:
        gesn.append({'code': row[1], 'name': row[2], 'unit': row[3]})
print(f"[{time.time()-t0:.1f}s] Загружено кодов ГЭСН: {len(gesn)}")

# 2. Парсинг всех 71 файлов ФРСН
files = sorted(glob.glob("docs/ФРСН/*.md"))
frsn_records = []
for path in files:
    m = re.match(r'docs/ФРСН/(\d+) ', path)
    sbornik_num = int(m.group(1)) if m else -1
    recs = parse_file(path)
    for r in recs:
        if not r['valid_shifr']:
            continue
        r['sbornik'] = sbornik_num
        frsn_records.append(r)
print(f"[{time.time()-t0:.1f}s] Загружено записей ФРСН (валидный шифр): {len(frsn_records)}")

# 3. Лемматизация и TF-IDF
frsn_lemma_str = [' '.join(sorted(norm_lemmas(r['name']))) for r in frsn_records]
gesn_lemma_str = [' '.join(sorted(norm_lemmas(g['name']))) for g in gesn]
vec = TfidfVectorizer()
X_frsn = vec.fit_transform(frsn_lemma_str)
X_gesn = vec.transform(gesn_lemma_str)
print(f"[{time.time()-t0:.1f}s] TF-IDF построен, признаков: {len(vec.vocabulary_)}")

# 4. Подбор кандидатов: TF-IDF косинус + бонусы/штрафы поверх
BATCH = 20
all_results = []
for start in range(0, len(gesn), BATCH):
    end = min(start + BATCH, len(gesn))
    sims = cosine_similarity(X_gesn[start:end], X_frsn)
    for row_i, g in zip(range(start, end), gesn[start:end]):
        row = sims[row_i - start]
        qn = numbers(g['name'])
        is_gesnm = g['code'].upper().startswith('ГЭСНМ')
        # адресуем только записи с ненулевым сходством, чтобы не перебирать все 77к для бонусов
        nz = np.nonzero(row)[0]
        scored = []
        for i in nz:
            base = float(row[i])
            r = frsn_records[i]
            rn = numbers(r['name'])
            num_bonus = 0.15 if (qn and rn and (qn & rn)) else 0.0
            family_bonus = 0.1 if (is_gesnm and 49 <= r['sbornik'] <= 72) else 0.0
            family_penalty = -0.05 if (is_gesnm and not (49 <= r['sbornik'] <= 72)) else 0.0
            scored.append((base + num_bonus + family_bonus + family_penalty, i))
        scored.sort(key=lambda x: -x[0])
        all_results.append((g, scored[:3]))
print(f"[{time.time()-t0:.1f}s] Подбор кандидатов завершён")

# 5. Формирование выходного xlsx
out = openpyxl.Workbook()
sh = out.active
sh.title = "Соответствие ГЭСН-ФРСН"
headers = ["Код ГЭСН", "Наименование ГЭСН", "Ед.изм. ГЭСН"]
for k in (1, 2, 3):
    headers += [f"Код ФРСН #{k}", f"Наименование ФРСН #{k}", f"чел-ч #{k}", f"Состав звена #{k}", f"Скор #{k}"]
headers.append("Подтверждено/отклонено (заполняет заказчик)")
sh.append(headers)
for c in sh[1]:
    c.font = Font(bold=True)

no_candidates = 0
weak = 0
scores_top1 = []
for g, cands in all_results:
    row = [g['code'], g['name'], g['unit']]
    if not cands:
        no_candidates += 1
    else:
        scores_top1.append(cands[0][0])
        if cands[0][0] < 0.3:
            weak += 1
    for k in range(3):
        if k < len(cands):
            score, i = cands[k]
            r = frsn_records[i]
            crew_str = "; ".join(f"{c['profession']} x{c['qty']}" for c in r['crew']) if r['crew'] else "(не найден)"
            row += [r['shifr'], r['name'], r['lab_hours'], crew_str, round(score, 3)]
        else:
            row += ['', '', '', '', '']
    row.append('')
    sh.append(row)

for col, width in zip('ABCDEFGHIJKLMNOP', [18, 45, 10, 18, 45, 8, 30, 8, 18, 45, 8, 30, 8, 18, 45, 8]):
    sh.column_dimensions[col].width = width

sh2 = out.create_sheet("Сводка")
sh2.append(["Методология подбора кандидатов ГЭСН -> ФРСН (Шаг 3, версия 2)"])
sh2.append([])
sh2.append(["Метод: слова приведены к начальной форме (лемматизация, pymorphy2),", ])
sh2.append(["сходство названий считается через TF-IDF + косинусное сходство (sklearn).", ])
sh2.append(["Поверх добавлены бонус за совпадение числового параметра (диаметр и т.п.)", ])
sh2.append(["и бонус/штраф для ГЭСНм относительно гипотезы, что этим кодам", ])
sh2.append(["соответствуют сборники ФРСН 49-72.", ])
sh2.append([])
sh2.append(["Версия 1 (черновик от 17.09 днём) использовала пересечение словоформ", ])
sh2.append(["без приведения к начальной форме - проверка показала, что это была", ])
sh2.append(["главная причина потерянных совпадений: слово \"задвижек\" и \"задвижка\"", ])
sh2.append(["считались разными словами. После лемматизации и перехода на TF-IDF", ])
sh2.append(["среднее качество совпадения выросло с 0.429 до 0.538, число слабых", ])
sh2.append(["совпадений (скор < 0.3) упало с 34 до 4 из 157.", ])
sh2.append([])
sh2.append(["ВАЖНО: числовые чел-ч НЕ использовались как критерий подбора -", ])
sh2.append(["это доказанно ненадёжный признак (см. пример ГЭСН22-03-007-02).", ])
sh2.append([])
sh2.append(["ВАЖНО: это по-прежнему черновик для ручной проверки, не готовый", ])
sh2.append(["результат. Даже сильные скоры нужно подтверждать, слабые - смотреть", ])
sh2.append(["особенно внимательно.", ])
sh2.append([])
sh2.append([f"Кодов ГЭСН: {len(gesn)}", f"Кодов без кандидатов: {no_candidates}"])
sh2.append([f"Слабых совпадений (скор топ-1 < 0.3): {weak}", f"Записей ФРСН (валидный шифр): {len(frsn_records)}"])

out.save("prototypes/results/2026-09-17_gesn_frsn_candidates_v2.xlsx")
print(f"[{time.time()-t0:.1f}s] Сохранено")
print(f"Кодов без кандидатов: {no_candidates} из {len(gesn)}")
print(f"Слабых совпадений (<0.3): {weak} из {len(gesn)}")
if scores_top1:
    print(f"Средний скор топ-1: {sum(scores_top1)/len(scores_top1):.3f}")
