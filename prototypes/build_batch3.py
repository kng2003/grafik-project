# -*- coding: utf-8 -*-
"""
Шаг 4 фазы 4: третья партия кодов ГЭСН на подтверждение заказчику -
коды со слабым скором (0.3-0.5). Партии 1 и 2 уже закрыты полностью.
На этом уровне скора совпадение по смыслу может быть верным, но часто
подобрано только по общим словам - нужна повышенная внимательность.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

SRC = "prototypes/results/2026-09-17_gesn_frsn_candidates_v2.xlsx"
OUT = "prototypes/results/2026-09-19_shag4_partiya3_slabye.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Соответствие ГЭСН-ФРСН"]
rows = list(ws.iter_rows(min_row=2, values_only=True))

def numscore(r):
    v = r[7]
    return v if isinstance(v, (int, float)) else None

selected = [r for r in rows if (numscore(r) is not None and 0.3 <= numscore(r) < 0.5)]
selected.sort(key=lambda r: numscore(r), reverse=True)

out = openpyxl.Workbook()
sh = out.active
sh.title = "Партия 3 - слабые"

headers = ["№", "Код ГЭСН", "Наименование ГЭСН", "Ед.изм.", "Код ФРСН (кандидат #1)",
           "Наименование ФРСН #1", "чел-ч #1", "Состав звена #1", "Скор #1",
           "Код ФРСН (кандидат #2)", "Наименование ФРСН #2", "чел-ч #2", "Состав звена #2", "Скор #2",
           "Ваше решение (да / нет / уточнить)", "Комментарий"]
sh.append(headers)
for c in sh[1]:
    c.font = Font(bold=True)
    c.alignment = Alignment(wrap_text=True, vertical="top")

yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

for i, r in enumerate(selected, start=1):
    # r columns: code,name,unit, [код1,назв1,челч1,звено1,скор1], [код2,назв2,челч2,звено2,скор2], [код3...], решение
    row = [i, r[0], r[1], r[2],
           r[3], r[4], r[5], r[6], r[7],
           r[8], r[9], r[10], r[11], r[12],
           "", ""]
    sh.append(row)

widths = [5, 18, 40, 9, 18, 40, 8, 26, 7, 18, 40, 8, 26, 7, 22, 30]
for col, w in zip("ABCDEFGHIJKLMNOP", widths):
    sh.column_dimensions[col].width = w

for row_idx in range(2, sh.max_row + 1):
    sh.cell(row_idx, 15).fill = yellow

sh.freeze_panes = "B2"

sh2 = out.create_sheet("Как заполнить")
instructions = [
    "ЧТО ЭТО ЗА ФАЙЛ",
    "Это третья партия из 4. В неё вошли 45 кодов ГЭСН со слабым скором совпадения",
    "(0.3-0.5). На этом уровне робот часто цепляется за общие слова в названии,",
    "а не за суть работы - поэтому здесь показаны сразу ДВА кандидата на каждый код,",
    "а не один, как в партиях 1 и 2. Возможно, ни один из двух не подходит - тогда",
    "нужно писать 'нет' и, если знаете, что должно быть на самом деле - в комментарии.",
    "",
    "ЧТО НУЖНО СДЕЛАТЬ",
    "Посмотреть на оба кандидата, сравнить с наименованием ГЭСН. Если один из двух",
    "подходит - написать в комментарии, какой (первый или второй), и 'да'. Если ни",
    "один не подходит - 'нет' и, по возможности, что искать. Если не уверены -",
    "'уточнить' с пояснением.",
    "",
    "НА ЧТО ОБРАТИТЬ ОСОБОЕ ВНИМАНИЕ",
    "На этом скоре чаще встречаются: работы из совсем другой области (сходство",
    "чисто по словам), или составные названия, где робот распознал только часть",
    "смысла. Не стесняйтесь писать 'нет' - это ожидаемо для слабых совпадений,",
    "не ошибка с вашей стороны.",
    "",
    "СКОЛЬКО ВРЕМЕНИ ЭТО ЗАЙМЁТ",
    "45 строк, но на каждую нужно больше внимания, чем в партиях 1-2. Прикидочно",
    "40-60 минут.",
    "",
    "ЧТО ДАЛЬШЕ",
    "После этой партии остаётся один код, ГЭСН08-01-008-05 (шпонка гидроизоляции),",
    "который робот вообще не смог сопоставить ни с ФРСН, ни с ЕНиР - по нему",
    "нужен отдельный разговор, не через таблицу.",
]
for line in instructions:
    sh2.append([line])
sh2.column_dimensions['A'].width = 95
for row_idx in range(1, len(instructions) + 1):
    sh2.cell(row_idx, 1).alignment = Alignment(wrap_text=True)
sh2['A1'].font = Font(bold=True, size=13)

out.save(OUT)
print(f"Сохранено: {OUT}")
print(f"Строк в партии 3: {len(selected)}")
